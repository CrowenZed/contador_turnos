#!/usr/bin/env python3
from __future__ import annotations

import argparse
import unicodedata
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Sequence

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MONTHS = {
    "enero": 1,
    "ene": 1,
    "febrero": 2,
    "feb": 2,
    "marzo": 3,
    "mar": 3,
    "abril": 4,
    "abr": 4,
    "mayo": 5,
    "may": 5,
    "junio": 6,
    "jun": 6,
    "julio": 7,
    "jul": 7,
    "agosto": 8,
    "ago": 8,
    "septiembre": 9,
    "setiembre": 9,
    "sept": 9,
    "sep": 9,
    "octubre": 10,
    "oct": 10,
    "noviembre": 11,
    "nov": 11,
    "diciembre": 12,
    "dic": 12,
}


@dataclass
class DayColumn:
    col: int
    month: int
    day: int


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = unicodedata.normalize("NFKD", text)
    return "".join(ch for ch in text if not unicodedata.combining(ch))


def parse_month(value: object) -> Optional[int]:
    token = normalize_text(value).replace(".", " ").replace("-", " ")
    parts = [p for p in token.split() if p]
    for part in parts:
        if part in MONTHS:
            return MONTHS[part]
    if token in MONTHS:
        return MONTHS[token]
    return None


def find_sheet_case_insensitive(wb, expected: str):
    lower = expected.lower()
    for name in wb.sheetnames:
        if name.lower() == lower:
            return wb[name]
    raise ValueError(f"No se encontró la hoja '{expected}'. Hojas disponibles: {wb.sheetnames}")


def detect_year(ws, default_year: int = 2026) -> int:
    # Requisito operativo: generar siempre para 2026 de forma fiable.
    if default_year == 2026:
        return 2026
    for r in range(1, min(12, ws.max_row) + 1):
        for c in range(1, min(30, ws.max_column) + 1):
            value = ws.cell(r, c).value
            if isinstance(value, int) and 2000 <= value <= 2100:
                return value
            if isinstance(value, str):
                digits = "".join(ch for ch in value if ch.isdigit())
                if len(digits) == 4:
                    y = int(digits)
                    if 2000 <= y <= 2100:
                        return y
    return default_year


def detect_day_row(ws) -> int:
    best_row = -1
    best_count = -1
    for r in range(1, min(25, ws.max_row) + 1):
        count = 0
        for c in range(1, ws.max_column + 1):
            value = ws.cell(r, c).value
            if isinstance(value, int) and 1 <= value <= 31:
                count += 1
        if count > best_count:
            best_count = count
            best_row = r
    if best_count < 7:
        raise ValueError("No pude detectar la fila de días del calendario (1..31).")
    return best_row


def detect_month_row(ws, day_row: int) -> int:
    best_row = day_row - 1
    best_count = -1
    for r in range(max(1, day_row - 5), day_row):
        count = 0
        for c in range(1, ws.max_column + 1):
            if parse_month(ws.cell(r, c).value):
                count += 1
        if count > best_count:
            best_count = count
            best_row = r
    return best_row


def extract_day_columns(ws, day_row: int, month_row: int) -> List[DayColumn]:
    month_by_col: Dict[int, int] = {}
    current_month: Optional[int] = None
    for c in range(1, ws.max_column + 1):
        mv = parse_month(ws.cell(month_row, c).value)
        if mv is not None:
            current_month = mv
        month_by_col[c] = current_month

    result: List[DayColumn] = []
    for c in range(1, ws.max_column + 1):
        day = ws.cell(day_row, c).value
        if not (isinstance(day, int) and 1 <= day <= 31):
            continue

        month = month_by_col.get(c)
        if month is None:
            for back in range(c, max(0, c - 12), -1):
                month = parse_month(ws.cell(month_row, back).value)
                if month is not None:
                    break
        if month is None:
            continue
        result.append(DayColumn(col=c, month=month, day=day))

    if len(result) < 7:
        raise ValueError("No pude extraer suficientes columnas de días con mes asociado.")
    return result


def split_weeks(day_cols: Sequence[DayColumn]) -> List[List[DayColumn]]:
    weeks: List[List[DayColumn]] = []
    block: List[DayColumn] = []
    for item in day_cols:
        block.append(item)
        if len(block) == 7:
            weeks.append(block)
            block = []
    if block:
        weeks.append(block)
    return weeks


def copy_column_style(ws, target_col: int, source_col: int, max_row: int) -> None:
    for r in range(1, max_row + 1):
        src = ws.cell(r, source_col)
        dst = ws.cell(r, target_col)
        dst._style = copy(src._style)
        if src.has_style:
            dst.number_format = src.number_format
            dst.font = copy(src.font)
            dst.fill = copy(src.fill)
            dst.border = copy(src.border)
            dst.alignment = copy(src.alignment)
            dst.protection = copy(src.protection)


def col_shift(original_col: int, insert_positions: Sequence[int], width: int = 2) -> int:
    return sum(width for p in insert_positions if p <= original_col)


INVALID_SHIFT_MARKERS = {
    "",
    "F",
    "V",
    "V/T",
    "T/V",
    "F/V",
    "V/F",
    "T",
    "-",
    "_",
}


def normalize_turn_value(value: object) -> str:
    text = "" if value is None else str(value)
    return text.strip().upper().replace(" ", "")


def is_real_shift(value: object) -> bool:
    normalized = normalize_turn_value(value)
    if normalized in INVALID_SHIFT_MARKERS:
        return False
    if set(normalized) <= {"-", "_"}:
        return False
    return True


def build_lookup_expr(turn_ref: str, date_expr: str, hours_col: str) -> str:
    return (
        "LET("
        f"raw,{turn_ref}&\"\","
        "t,SUBSTITUTE(UPPER(TRIM(raw)),\" \",\"\"),"
        f"d,{date_expr},"
        "tipo,IF(COUNTIF(festivos!$A:$A,d)>0,\"FES\","
        "IF(WEEKDAY(d,2)=6,\"DIS\",IF(WEEKDAY(d,2)=7,\"FES\",\"LAB\"))),"
        "invalido,OR(t=\"\",t=\"F\",t=\"V\",t=\"V/T\",t=\"T/V\",t=\"F/V\",t=\"V/F\",t=\"T\",t=\"-\",t=\"_\"),"
        "IF(invalido,0,"
        "IF(tipo=\"LAB\",IFERROR(INDEX(LAB!$"
        f"{hours_col}:${hours_col},MATCH(--t,LAB!$A:$A,0)),0),"
        "IF(tipo=\"DIS\",IFERROR(INDEX(DIS!$"
        f"{hours_col}:${hours_col},MATCH(--t,DIS!$A:$A,0)),0),"
        "IFERROR(INDEX(FES!$"
        f"{hours_col}:${hours_col},MATCH(--t,FES!$A:$A,0)),0)"
        ")"
        ")"
        ")"
        ")"
    )


def detect_data_rows(ws, day_cols: Sequence[DayColumn], start_row: int) -> List[int]:
    rows: List[int] = []
    first_day_col = day_cols[0].col
    day_col_set = {d.col for d in day_cols}

    id_end = max(1, first_day_col - 1)
    for r in range(start_row, ws.max_row + 1):
        has_id = any(ws.cell(r, c).value not in (None, "") for c in range(1, min(id_end, 5) + 1))
        has_calendar_content = any(ws.cell(r, c).value not in (None, "") for c in day_col_set)
        has_real_turn = any(is_real_shift(ws.cell(r, c).value) for c in day_col_set)
        if has_id and (has_real_turn or has_calendar_content):
            rows.append(r)
    return rows


def prepare_workbook(input_path: Path, output_path: Path) -> None:
    wb = load_workbook(input_path)

    ws = find_sheet_case_insensitive(wb, "GENERAL 26")
    _ = find_sheet_case_insensitive(wb, "LAB")
    _ = find_sheet_case_insensitive(wb, "DIS")
    _ = find_sheet_case_insensitive(wb, "FES")
    _ = find_sheet_case_insensitive(wb, "festivos")

    year = detect_year(ws, default_year=2026)
    day_row = detect_day_row(ws)
    month_row = detect_month_row(ws, day_row)
    day_cols = extract_day_columns(ws, day_row, month_row)
    weeks = split_weeks(day_cols)

    max_row_before = ws.max_row
    insert_positions: List[int] = []

    for week in reversed(weeks):
        insert_at = week[-1].col + 1
        ws.insert_cols(insert_at, amount=2)
        copy_column_style(ws, insert_at, max(1, insert_at - 1), max_row_before)
        copy_column_style(ws, insert_at + 1, max(1, insert_at - 1), max_row_before)
        insert_positions.append(insert_at)

    shifted_days = [DayColumn(col=d.col + col_shift(d.col, insert_positions), month=d.month, day=d.day) for d in day_cols]
    shifted_weeks = split_weeks(shifted_days)

    data_rows = detect_data_rows(ws, shifted_days, day_row + 1)

    week_morning_cols: List[int] = []
    week_afternoon_cols: List[int] = []

    for week in shifted_weeks:
        morning_col = week[-1].col + 1
        afternoon_col = week[-1].col + 2
        week_morning_cols.append(morning_col)
        week_afternoon_cols.append(afternoon_col)

        ws.cell(day_row, morning_col, "Horas mañana")
        ws.cell(day_row, afternoon_col, "Horas tarde")

        for r in data_rows:
            morning_terms = []
            afternoon_terms = []
            for d in week:
                day_letter = get_column_letter(d.col)
                turn_ref = f"{day_letter}{r}"
                date_expr = f"DATE({int(year)},{d.month},{d.day})"
                morning_terms.append(build_lookup_expr(turn_ref, date_expr, "B"))
                afternoon_terms.append(build_lookup_expr(turn_ref, date_expr, "C"))

            ws.cell(r, morning_col, f"=SUM({','.join(morning_terms)})")
            ws.cell(r, afternoon_col, f"=SUM({','.join(afternoon_terms)})")
            ws.cell(r, morning_col).number_format = "[h]:mm"
            ws.cell(r, afternoon_col).number_format = "[h]:mm"

    total_insert_at = ws.max_column + 1
    ws.insert_cols(total_insert_at, amount=2)
    copy_column_style(ws, total_insert_at, max(1, total_insert_at - 1), ws.max_row)
    copy_column_style(ws, total_insert_at + 1, max(1, total_insert_at - 1), ws.max_row)

    ws.cell(day_row, total_insert_at, "Total mañana")
    ws.cell(day_row, total_insert_at + 1, "Total tarde")

    for r in data_rows:
        morning_refs = [f"{get_column_letter(c)}{r}" for c in week_morning_cols]
        afternoon_refs = [f"{get_column_letter(c)}{r}" for c in week_afternoon_cols]

        ws.cell(r, total_insert_at, f"=SUM({','.join(morning_refs)})")
        ws.cell(r, total_insert_at + 1, f"=SUM({','.join(afternoon_refs)})")
        ws.cell(r, total_insert_at).number_format = "[h]:mm"
        ws.cell(r, total_insert_at + 1).number_format = "[h]:mm"

    wb.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Prepara VACACIONES.xlsx con columnas y fórmulas automáticas.")
    parser.add_argument("--input", default="VACACIONES.xlsx", help="Archivo Excel de entrada")
    parser.add_argument("--output", default="VACACIONES_preparado.xlsx", help="Archivo Excel de salida")
    args = parser.parse_args()

    prepare_workbook(Path(args.input), Path(args.output))
    print(f"Archivo generado: {args.output}")


if __name__ == "__main__":
    main()
